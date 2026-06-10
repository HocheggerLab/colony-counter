"""Excel and CSV report writers."""

from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.worksheet import Worksheet

from .detection import DetectionResult


@dataclass
class ImageReport:
    """Combines a source image's name with its detection result."""

    name: str
    result: DetectionResult


# ── Reusable styles ─────────────────────────────────────────────────────────
HEADER_FONT = Font(name="Arial", bold=True, color="FFFFFF", size=11)
HEADER_FILL = PatternFill("solid", start_color="2F4F7F")
SUBHEADER_FILL = PatternFill("solid", start_color="D9E1F2")
SUBHEADER_FONT = Font(name="Arial", bold=True, size=10)
CELL_FONT = Font(name="Arial", size=10)
TITLE_FONT = Font(name="Arial", bold=True, size=13, color="2F4F7F")
CENTER = Alignment(horizontal="center", vertical="center")
LEFT = Alignment(horizontal="left", vertical="center")
THIN = Side(style="thin", color="BBBBBB")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
EVEN_FILL = PatternFill("solid", start_color="F2F5FB")
ODD_FILL = PatternFill("solid", start_color="FFFFFF")

SUMMARY_HEADERS = [
    "Image Name",
    "Colony Count",
    "Colony Area (%)",
    "Mean Density (0-100)",
    "Dish Area (px)",
    "Mask Image",
]
COLONY_HEADERS = [
    "Colony ID",
    "Area (px)",
    "Area (%)",
    "Density (0-100)",
    "Centroid X (px)",
    "Centroid Y (px)",
]
COLONY_COL_WIDTHS = [12, 12, 12, 18, 18, 18]


def _style_header_row(ws: Worksheet, row: int, headers: list[str]) -> None:
    for col, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=col, value=h)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = CENTER
        cell.border = BORDER


def _write_summary_sheet(ws: Worksheet, reports: list[ImageReport]) -> None:
    ws.title = "Summary"
    ws.merge_cells("A1:F1")
    ws["A1"] = "Colony Survival Assay — Results Summary"
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = LEFT
    ws.row_dimensions[1].height = 28
    ws.row_dimensions[2].height = 6

    _style_header_row(ws, 3, SUMMARY_HEADERS)
    ws.row_dimensions[3].height = 20

    for i, report in enumerate(reports, start=4):
        r = report.result
        row_fill = EVEN_FILL if i % 2 == 0 else ODD_FILL
        values: list[object] = [
            report.name,
            r.n_colonies,
            round(r.colony_area_pct, 1),
            round(r.mean_density, 1),
            r.dish_area_px,
            f"{Path(report.name).stem}_mask.png",
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.font = CELL_FONT
            cell.fill = row_fill
            cell.border = BORDER
            cell.alignment = CENTER if col > 1 else LEFT
        ws.row_dimensions[i].height = 18

    n_rows = len(reports)
    summary_row = 4 + n_rows
    ws.row_dimensions[summary_row].height = 20
    for col in range(1, 7):
        cell = ws.cell(row=summary_row, column=col)
        cell.fill = SUBHEADER_FILL
        cell.border = BORDER
        cell.font = SUBHEADER_FONT
        cell.alignment = CENTER
    ws.cell(row=summary_row, column=1, value="AVERAGE")
    if n_rows:
        for col_letter, col_idx in (("B", 2), ("C", 3), ("D", 4)):
            ws.cell(
                row=summary_row,
                column=col_idx,
                value=f"=AVERAGE({col_letter}4:{col_letter}{summary_row - 1})",
            ).number_format = "0.0"

        for row in ws.iter_rows(
            min_row=4, max_row=summary_row - 1, min_col=3, max_col=4
        ):
            for cell in row:
                cell.number_format = "0.0"

    widths = {"A": 40, "B": 16, "C": 18, "D": 22, "E": 18, "F": 38}
    for letter, width in widths.items():
        ws.column_dimensions[letter].width = width


def _write_colony_sheet(ws: Worksheet, report: ImageReport) -> None:
    r = report.result
    ws.merge_cells("A1:F1")
    ws["A1"] = (
        f"{report.name}  —  {r.n_colonies} colonies  |  "
        f"Area: {r.colony_area_pct:.1f}%  |  "
        f"Mean density: {r.mean_density:.1f}"
    )
    ws["A1"].font = TITLE_FONT
    ws["A1"].alignment = LEFT
    ws.row_dimensions[1].height = 26
    ws.row_dimensions[2].height = 6

    _style_header_row(ws, 3, COLONY_HEADERS)
    ws.row_dimensions[3].height = 20

    for i, c in enumerate(r.colonies, start=4):
        row_fill = EVEN_FILL if i % 2 == 0 else ODD_FILL
        values: list[object] = [
            c["colony_id"],
            c["area_px"],
            c["area_pct"],
            c["density"],
            c["centroid_x"],
            c["centroid_y"],
        ]
        for col, val in enumerate(values, start=1):
            cell = ws.cell(row=i, column=col, value=val)
            cell.font = CELL_FONT
            cell.fill = row_fill
            cell.border = BORDER
            cell.alignment = CENTER
        ws.row_dimensions[i].height = 16

    last_data = 3 + len(r.colonies)
    avg_row = last_data + 1
    ws.row_dimensions[avg_row].height = 20
    for col in range(1, 7):
        cell = ws.cell(row=avg_row, column=col)
        cell.fill = SUBHEADER_FILL
        cell.border = BORDER
        cell.font = SUBHEADER_FONT
        cell.alignment = CENTER
    ws.cell(row=avg_row, column=1, value="AVERAGE")
    if r.colonies:
        ws.cell(
            row=avg_row,
            column=2,
            value=f"=AVERAGE(B4:B{last_data})",
        ).number_format = "0"
        ws.cell(
            row=avg_row,
            column=3,
            value=f"=AVERAGE(C4:C{last_data})",
        ).number_format = "0.000"
        ws.cell(
            row=avg_row,
            column=4,
            value=f"=AVERAGE(D4:D{last_data})",
        ).number_format = "0.00"

    for width, letter in zip(
        COLONY_COL_WIDTHS, ["A", "B", "C", "D", "E", "F"], strict=True
    ):
        ws.column_dimensions[letter].width = width


def write_csvs(reports: list[ImageReport], output_dir: Path) -> None:
    """Write summary.csv and colonies.csv (long format) to output_dir."""
    output_dir.mkdir(parents=True, exist_ok=True)

    with open(output_dir / "summary.csv", "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(
            ["image_name", "colony_count", "colony_area_pct", "mean_density", "dish_area_px"]
        )
        for report in reports:
            r = report.result
            writer.writerow(
                [report.name, r.n_colonies, r.colony_area_pct, r.mean_density, r.dish_area_px]
            )

    with open(output_dir / "colonies.csv", "w", newline="") as f:
        writer = csv.writer(f)
        writer.writerow(
            ["image_name", "colony_id", "area_px", "area_pct", "density", "centroid_x", "centroid_y"]
        )
        for report in reports:
            for c in report.result.colonies:
                writer.writerow(
                    [
                        report.name,
                        c["colony_id"],
                        c["area_px"],
                        c["area_pct"],
                        c["density"],
                        c["centroid_x"],
                        c["centroid_y"],
                    ]
                )


def write_workbook(reports: list[ImageReport], output_path: Path) -> None:
    """Write a styled Excel workbook with a summary and per-image sheets."""
    wb = Workbook()
    summary_ws = wb.active
    if summary_ws is None:  # pragma: no cover - openpyxl always creates one
        summary_ws = wb.create_sheet()
    _write_summary_sheet(summary_ws, reports)

    for report in reports:
        sheet_name = Path(report.name).stem[:31]
        ws = wb.create_sheet(title=sheet_name)
        _write_colony_sheet(ws, report)

    output_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(output_path)
