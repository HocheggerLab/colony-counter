"""Tests for the Excel report writer."""

from __future__ import annotations

from pathlib import Path

import numpy as np
from openpyxl import load_workbook

import csv

from colony_counter.detection import ColonyRecord, DetectionResult
from colony_counter.excel import ImageReport, write_csvs, write_workbook


def _fake_result(n: int = 2) -> DetectionResult:
    colonies: list[ColonyRecord] = [
        ColonyRecord(
            colony_id=i + 1,
            area_px=100 + i * 10,
            area_pct=1.0 + i * 0.1,
            density=50.0 + i,
            centroid_x=10.0 * (i + 1),
            centroid_y=20.0 * (i + 1),
        )
        for i in range(n)
    ]
    return DetectionResult(
        image=np.zeros((4, 4, 3), dtype=np.uint8),
        labels=np.zeros((4, 4), dtype=int),
        colonies=colonies,
        n_colonies=n,
        dish_area_px=10_000,
        colony_area_px=sum(c["area_px"] for c in colonies),
        colony_area_pct=2.1,
        mean_density=50.5,
    )


def test_write_workbook_creates_summary_and_per_image_sheets(
    tmp_path: Path,
) -> None:
    reports = [
        ImageReport(name="plate_01.tif", result=_fake_result(2)),
        ImageReport(name="plate_02.tif", result=_fake_result(3)),
    ]
    out = tmp_path / "results.xlsx"
    write_workbook(reports, out)

    assert out.exists()
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary", "plate_01", "plate_02"]

    summary = wb["Summary"]
    # Header row at row 3
    assert summary.cell(row=3, column=1).value == "Image Name"
    # First data row
    assert summary.cell(row=4, column=1).value == "plate_01.tif"
    assert summary.cell(row=4, column=2).value == 2
    # AVERAGE row follows last data row
    assert summary.cell(row=6, column=1).value == "AVERAGE"


def test_write_workbook_with_empty_reports(tmp_path: Path) -> None:
    out = tmp_path / "empty.xlsx"
    write_workbook([], out)
    wb = load_workbook(out)
    assert wb.sheetnames == ["Summary"]


def test_write_csvs_summary_rows(tmp_path: Path) -> None:
    reports = [
        ImageReport(name="plate_01.tif", result=_fake_result(2)),
        ImageReport(name="plate_02.tif", result=_fake_result(1)),
    ]
    write_csvs(reports, tmp_path)

    summary = list(csv.DictReader((tmp_path / "summary.csv").open()))
    assert len(summary) == 2
    assert summary[0]["image_name"] == "plate_01.tif"
    assert int(summary[0]["colony_count"]) == 2
    assert summary[1]["image_name"] == "plate_02.tif"


def test_write_csvs_colonies_long_format(tmp_path: Path) -> None:
    reports = [
        ImageReport(name="plate_01.tif", result=_fake_result(2)),
        ImageReport(name="plate_02.tif", result=_fake_result(3)),
    ]
    write_csvs(reports, tmp_path)

    colonies = list(csv.DictReader((tmp_path / "colonies.csv").open()))
    # 2 + 3 colonies total
    assert len(colonies) == 5
    assert colonies[0]["image_name"] == "plate_01.tif"
    assert colonies[2]["image_name"] == "plate_02.tif"
    assert set(colonies[0].keys()) == {
        "image_name", "colony_id", "area_px", "area_pct",
        "density", "centroid_x", "centroid_y",
    }


def test_write_csvs_empty_reports(tmp_path: Path) -> None:
    write_csvs([], tmp_path)
    assert (tmp_path / "summary.csv").exists()
    assert (tmp_path / "colonies.csv").exists()
    assert list(csv.DictReader((tmp_path / "summary.csv").open())) == []


def test_long_image_name_is_truncated_for_sheet(tmp_path: Path) -> None:
    long_name = "a" * 50 + ".tif"
    reports = [ImageReport(name=long_name, result=_fake_result(1))]
    out = tmp_path / "long.xlsx"
    write_workbook(reports, out)
    wb = load_workbook(out)
    # Excel sheet name limit is 31 chars
    assert len(wb.sheetnames[1]) <= 31
